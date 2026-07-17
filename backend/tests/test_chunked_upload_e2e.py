"""
E2E tests for the chunked upload flow (calamine parser + heartbeat + backup chunking).
Iteration 32 - covers: happy path, backup chunking + undo, heartbeat stale detection,
unknown upload_id 404, and 'no valid records' error path.

Run: pytest /app/backend/tests/test_chunked_upload_e2e.py -v -s
"""
import os, io, math, time, random, datetime, pytest, requests
from pymongo import MongoClient
from openpyxl import Workbook

BASE = os.environ.get("REACT_APP_BACKEND_URL", "https://alias-mapping-debug.preview.emergentagent.com").rstrip("/")
CHUNK = 200 * 1024
MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/api/auth/login", json={"username": "admin", "password": "admin123"}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def H(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def db():
    return MongoClient(MONGO_URL)[DB_NAME]


def _make_small_sale(path, n=1500, dstart=(2025, 1, 1)):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["Some Title Row"])
    ws.append(["Date", "Type", "Refno", "Party Name", "Item Name", "Stamp", "Lbr. On Tag.No.",
               "On", "Gr.Wt.", "Gold Std.", "Fine", "Tunch", "Total", "Taxable Val.", "Pc"])
    random.seed(1)
    base = datetime.date(*dstart)
    for i in range(n):
        d = base + datetime.timedelta(days=random.randint(0, 300))
        ws.append([d, "S", f"R{i}", f"PARTY {i%50}", "CHAIN X", "S 70", "T-1",
                   round(random.uniform(5, 30), 2), round(random.uniform(0.1, 5), 3),
                   round(random.uniform(0.1, 5), 3), round(random.uniform(0.05, 4), 3),
                   round(random.uniform(60, 90), 2), round(random.uniform(1000, 9000), 2),
                   round(random.uniform(900, 8000), 2), random.randint(1, 20)])
    wb.save(path)


def _make_invalid_sale(path):
    """Rows with 1-char item names -> parser rejects all -> 'no valid records' error."""
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["Some Title Row"])
    ws.append(["Date", "Type", "Refno", "Party Name", "Item Name", "Stamp", "Lbr. On Tag.No.",
               "On", "Gr.Wt.", "Gold Std.", "Fine", "Tunch", "Total", "Taxable Val.", "Pc"])
    # >200KB threshold: pad with many junk rows (single char item -> rejected)
    for i in range(15000):
        ws.append([datetime.date(2025, 6, 1), "S", f"R{i}", f"P{i%20}", "X",
                   "S 70", "T-1", 10, 1, 1, 0.5, 75, 1000, 900, 1])
    wb.save(path)


def _upload_file(path, H, start="2025-01-01", end="2025-12-31", file_type="sale"):
    size = os.path.getsize(path)
    total_chunks = math.ceil(size / CHUNK)
    r = requests.post(f"{BASE}/api/upload/init",
                      json={"file_type": file_type, "start_date": start, "end_date": end,
                            "total_chunks": total_chunks}, headers=H, timeout=30)
    assert r.status_code == 200, r.text
    uid = r.json()["upload_id"]
    with open(path, "rb") as f:
        for i in range(total_chunks):
            data = f.read(CHUNK)
            rr = requests.post(f"{BASE}/api/upload/chunk/{uid}?chunk_index={i}",
                               files={"file": (f"c{i}", data)}, headers=H, timeout=60)
            assert rr.status_code == 200, rr.text
    r = requests.post(f"{BASE}/api/upload/finalize/{uid}", headers=H, timeout=30)
    assert r.status_code == 200, r.text
    return uid, size, total_chunks


def _poll(uid, H, timeout_s=300):
    t0 = time.time()
    last = None
    while time.time() - t0 < timeout_s:
        r = requests.get(f"{BASE}/api/upload/status/{uid}", headers=H, timeout=30)
        if r.status_code != 200:
            return {"status": "http_error", "code": r.status_code, "body": r.text, "elapsed": time.time() - t0}
        s = r.json()
        last = s
        st = s.get("status")
        if st in ("complete", "error", "failed"):
            s["elapsed"] = time.time() - t0
            return s
        time.sleep(2)
    last["elapsed"] = time.time() - t0
    last["status"] = "timeout"
    return last


# ---- Test 1: happy path (small sale, chunked) ----
def test_chunked_upload_happy_path(H, db):
    path = "/tmp/small_sale.xlsx"
    _make_small_sale(path, n=1500)
    assert os.path.getsize(path) > 25 * 1024  # non-trivial
    uid, size, chunks = _upload_file(path, H)
    print(f"[happy] size={size/1024:.1f}KB chunks={chunks} uid={uid}")
    s = _poll(uid, H, timeout_s=180)
    print(f"[happy] final status={s.get('status')} elapsed={s.get('elapsed'):.1f}s msg={s.get('message')}")
    assert s["status"] == "complete", s
    assert s["elapsed"] < 120, f"took too long: {s['elapsed']:.1f}s"
    batch_id = s.get("batch_id")
    assert batch_id
    cnt = db.transactions.count_documents({"batch_id": batch_id})
    assert cnt == 1500, cnt
    # cleanup
    db.transactions.delete_many({"batch_id": batch_id})
    db.replaced_records.delete_many({"batch_id": batch_id})
    db.action_history.delete_many({"batch_id": batch_id})


# ---- Test 2: backup chunking + undo ----
def test_backup_chunking_and_undo(H, db):
    # Use existing big_sale.xlsx (60K rows) - upload twice
    path = "/tmp/big_sale.xlsx"
    if not os.path.exists(path):
        pytest.skip("big_sale.xlsx not present; run /tmp/test_parse.py to generate")
    # Clean any existing 2025 data first to isolate
    db.transactions.delete_many({"type": {"$in": ["sale", "sale_return"]},
                                 "date": {"$gte": "2025-01-01", "$lte": "2025-12-31"}})
    # First upload
    uid1, _, _ = _upload_file(path, H)
    s1 = _poll(uid1, H, timeout_s=300)
    print(f"[backup] first upload elapsed={s1.get('elapsed'):.1f}s status={s1.get('status')}")
    assert s1["status"] == "complete", s1
    batch1 = s1["batch_id"]
    cnt1 = db.transactions.count_documents({"batch_id": batch1})
    assert cnt1 == 60000, cnt1

    # Second upload (replaces)
    uid2, _, _ = _upload_file(path, H)
    s2 = _poll(uid2, H, timeout_s=300)
    print(f"[backup] second upload elapsed={s2.get('elapsed'):.1f}s status={s2.get('status')}")
    assert s2["status"] == "complete", s2
    batch2 = s2["batch_id"]

    parts = list(db.replaced_records.find({"batch_id": batch2}))
    print(f"[backup] replaced_records parts={len(parts)}")
    assert len(parts) >= 2, f"expected multiple chunks, got {len(parts)}"
    for p in parts:
        assert "part" in p, p.keys()
        # each part must contain <=5000 records
        recs = p.get("records") or p.get("data") or []
        assert len(recs) <= 5000, len(recs)

    # Undo second batch -> should restore first batch's records
    r = requests.post(f"{BASE}/api/history/undo-upload?batch_id={batch2}", headers=H, timeout=120)
    print(f"[backup] undo status={r.status_code} body={r.text[:200]}")
    assert r.status_code == 200, r.text

    restored = db.transactions.count_documents({"type": "sale",
                                                 "date": {"$gte": "2025-01-01", "$lte": "2025-12-31"}})
    print(f"[backup] restored count={restored}")
    assert restored == 60000, restored

    # Cleanup
    db.transactions.delete_many({"type": {"$in": ["sale", "sale_return"]},
                                 "date": {"$gte": "2025-01-01", "$lte": "2025-12-31"}})
    db.replaced_records.delete_many({"batch_id": {"$in": [batch1, batch2]}})
    db.action_history.delete_many({"batch_id": {"$in": [batch1, batch2]}})


# ---- Test 3: heartbeat stale detection ----
def test_heartbeat_stale(H, db):
    fake_uid = f"test_stale_{int(time.time())}"
    # heartbeat is stored as ISO-8601 string (see _save_upload_meta in server.py)
    stale_dt = datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(minutes=5)
    stale_ts = stale_dt.isoformat()
    doc = {
        "upload_id": fake_uid,
        "status": "processing",
        "owner_username": "admin",
        "heartbeat": stale_ts,
        "file_type": "sale",
        "start_date": "2025-01-01",
        "end_date": "2025-12-31",
        "created_at": stale_ts,
    }
    db.upload_sessions.insert_one(doc)
    try:
        r = requests.get(f"{BASE}/api/upload/status/{fake_uid}", headers=H, timeout=30)
        print(f"[stale] status_code={r.status_code} body={r.text[:400]}")
        assert r.status_code == 200, r.text
        s = r.json()
        assert s["status"] == "error", s
        blob = ((s.get("detail") or "") + " " + (s.get("message") or "")).lower()
        assert "stop" in blob or "unexpected" in blob or "stale" in blob or "timed out" in blob, s
    finally:
        db.upload_sessions.delete_many({"upload_id": fake_uid})


# ---- Test 4: unknown id -> 404 ----
def test_unknown_upload_id_404(H):
    r = requests.get(f"{BASE}/api/upload/status/nonexistent_xyz_123", headers=H, timeout=30)
    print(f"[404] status_code={r.status_code} body={r.text[:200]}")
    assert r.status_code == 404, r.text


# ---- Test 5: no valid records path ----
def test_no_valid_records_error(H, db):
    path = "/tmp/invalid_sale.xlsx"
    _make_invalid_sale(path)
    size = os.path.getsize(path)
    print(f"[invalid] file size = {size/1024:.1f}KB")
    assert size > 200 * 1024, f"file too small ({size}) - won't hit chunked path"
    uid, _, _ = _upload_file(path, H)
    s = _poll(uid, H, timeout_s=180)
    print(f"[invalid] final status={s.get('status')} detail={s.get('detail')} msg={s.get('message')}")
    assert s["status"] == "error", s
    blob = ((s.get("detail") or "") + " " + (s.get("message") or "")).lower()
    assert "no valid" in blob or "no records" in blob or "valid records" in blob, s


# ---- Test 6: /api/stats regression ----
def test_stats_regression(H):
    r = requests.get(f"{BASE}/api/stats", headers=H, timeout=30)
    assert r.status_code == 200, r.text
    s = r.json()
    print(f"[stats] total_transactions={s.get('total_transactions')}")
    # Note: 2025-dated rows may transiently exist during the run; snapshot at end after cleanup
    assert s.get("total_transactions", 0) >= 11767
