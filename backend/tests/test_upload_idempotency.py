"""
Regression tests for the RE-UPLOAD IDEMPOTENCY bug.

Root cause: Tally continuation lines (rows with item+weights but empty
date/refno/party/type) were inserted with date='' and never replaced on
re-upload (the replace-scope filter used only new-file dates). Re-uploads
therefore DUPLICATED ghost rows and mis-stated stock.

Fixes tested:
  1. Forward-fill of date/refno/party/type in sale & purchase mappers.
  2. Replace scope now includes date IN new_dates OR '' / None -> ghost
     rows are purged on the next same-type upload (self-repair).
  3. Numeric grand-total rows are NOT inserted (parser drops rows with
     no valid item name).
  4. sale_return (Type='R') inheritance to continuation lines.
"""
import datetime
import math
import os
import random
import time
import uuid

import pytest
import requests
from openpyxl import Workbook
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "http://localhost:8001").rstrip("/")
MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")

db = MongoClient(MONGO_URL)[DB_NAME]

TAG = "TEST_IDEMPOT_" + uuid.uuid4().hex[:8]

# ---------- helpers ----------

def _auth_token():
    r = requests.post(
        f"{BASE_URL}/api/auth/login",
        json={"username": "admin", "password": "admin123"},
        timeout=15,
    )
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def headers():
    return {"Authorization": f"Bearer {_auth_token()}"}


def _upload_chunked(path, file_type, headers, start="2025-01-01", end="2025-12-31"):
    size = os.path.getsize(path)
    CHUNK = 256 * 1024
    n = max(1, math.ceil(size / CHUNK))
    r = requests.post(
        f"{BASE_URL}/api/upload/init",
        json={
            "file_type": file_type,
            "start_date": start,
            "end_date": end,
            "total_chunks": n,
        },
        headers=headers,
        timeout=30,
    )
    assert r.status_code == 200, r.text
    uid = r.json()["upload_id"]
    with open(path, "rb") as f:
        for i in range(n):
            chunk = f.read(CHUNK)
            rr = requests.post(
                f"{BASE_URL}/api/upload/chunk/{uid}?chunk_index={i}",
                files={"file": (f"c{i}", chunk)},
                headers=headers,
                timeout=60,
            )
            assert rr.status_code == 200, rr.text
    fr = requests.post(f"{BASE_URL}/api/upload/finalize/{uid}", headers=headers, timeout=60)
    assert fr.status_code == 200, fr.text
    # poll
    for _ in range(180):  # up to 9 min
        time.sleep(3)
        s = requests.get(f"{BASE_URL}/api/upload/status/{uid}", headers=headers, timeout=15).json()
        if s.get("status") in ("complete", "error"):
            return s
    raise TimeoutError("upload did not finish")


def _make_tally_sale(path, n_vouchers=3000, seed=11, sale_return_voucher_idx=7):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["Some Title Row"])
    ws.append(
        [
            "Date", "Type", "Refno", "Party Name", "Item Name", "Stamp",
            "Lbr. On Tag.No.", "On", "Gr.Wt.", "Gold Std.", "Fine", "Tunch",
            "Total", "Taxable Val.", "Pc",
        ]
    )
    items = ["CHAIN MS-70", "PAYAL KJN", "SNT 40-256", "CHAIN KJN-70", "RING X"]
    random.seed(seed)
    total_rows = 0
    for v in range(n_vouchers):
        d = datetime.date(2025, 1, 1) + datetime.timedelta(days=random.randint(0, 364))
        # Force >=2 lines for the sale_return voucher so we can validate inheritance
        if v == sale_return_voucher_idx:
            n_lines, vtype = 3, "R"
        else:
            n_lines = random.choice([1, 1, 2, 3])
            vtype = "S"
        refno = f"V{v}"
        for line in range(n_lines):
            first = line == 0
            ws.append(
                [
                    d if first else None,
                    vtype if first else None,
                    refno if first else None,
                    f"PARTY {random.randint(0, 100)}" if first else None,
                    random.choice(items),
                    "STAMP 70",
                    "T-12",
                    round(random.uniform(5, 30), 2),
                    round(random.uniform(0.1, 5), 3),
                    round(random.uniform(0.1, 5), 3),
                    round(random.uniform(0.05, 4), 3),
                    round(random.uniform(60, 90), 2),
                    round(random.uniform(1000, 90000), 2),
                    round(random.uniform(900, 80000), 2),
                    random.randint(1, 20),
                ]
            )
            total_rows += 1
    # numeric grand-total row (all text cells are numbers) -> must be dropped
    ws.append(
        [
            None, None, 210226.0, 210226.0, 210226.0, 210226.0, None, None,
            43555.4, 39000.2, 20000.1, 12013453.5, 3.29e8, 3.1e8, 210226,
        ]
    )
    wb.save(path)
    return total_rows, f"V{sale_return_voucher_idx}"


def _make_small_purchase(path):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["Purchase Register"])
    ws.append(
        [
            "Date", "Type", "Refno", "Party Name", "Item Name", "Stamp",
            "Lbr. On Tag.No.", "On", "Gr.Wt.", "Gold Std.", "Fine", "Tunch",
            "Total", "Taxable Val.", "Pc",
        ]
    )
    random.seed(3)
    rows = 0
    for v in range(60):
        d = datetime.date(2025, 3, 1) + datetime.timedelta(days=v % 30)
        # first line dated, then a continuation line for ~half of them
        n_lines = 2 if v % 2 == 0 else 1
        for line in range(n_lines):
            first = line == 0
            ws.append(
                [
                    d if first else None,
                    "P" if first else None,
                    f"P{v}" if first else None,
                    f"SUPPLIER {v % 10}" if first else None,
                    "PURCHASE ITEM A",
                    "S70", "T-1",
                    round(random.uniform(5, 20), 2),
                    round(random.uniform(0.1, 3), 3),
                    round(random.uniform(0.1, 3), 3),
                    round(random.uniform(0.05, 2), 3),
                    round(random.uniform(60, 90), 2),
                    round(random.uniform(500, 20000), 2),
                    round(random.uniform(500, 20000), 2),
                    random.randint(1, 10),
                ]
            )
            rows += 1
    wb.save(path)
    return rows


def _sale_stats(batch_id):
    n_all = db.transactions.count_documents({"batch_id": batch_id})
    n_nodate = db.transactions.count_documents({"batch_id": batch_id, "date": ""})
    pipe = [
        {"$match": {"batch_id": batch_id, "type": "sale"}},
        {"$group": {"_id": None, "net": {"$sum": "$net_wt"}, "gross": {"$sum": "$gross_wt"}}},
    ]
    agg = list(db.transactions.aggregate(pipe))
    net = round(agg[0]["net"], 3) if agg else 0.0
    return n_all, n_nodate, net


def _cleanup_batches(batch_ids):
    if not batch_ids:
        return
    db.transactions.delete_many({"batch_id": {"$in": batch_ids}})
    db.replaced_records.delete_many({"batch_id": {"$in": batch_ids}})
    db.action_history.delete_many({"description": {"$regex": "records for 2025"}})


# ---------- tests ----------

BATCHES_TO_CLEAN = []


@pytest.fixture(scope="module", autouse=True)
def _final_cleanup():
    baseline = db.transactions.count_documents({})
    yield
    _cleanup_batches(BATCHES_TO_CLEAN)
    # Purge any stray legacy ghosts left by this test module
    db.transactions.delete_many({"batch_id": "legacy-ghosts-" + TAG})
    after = db.transactions.count_documents({})
    print(f"\n[cleanup] baseline={baseline} after={after}")


class TestIdempotencySaleUpload:
    """User's bug: same file uploaded twice must produce identical DB state."""

    path = "/tmp/idempot_sale.xlsx"

    @classmethod
    def setup_class(cls):
        rows, cls.sr_refno = _make_tally_sale(cls.path)
        cls.rows_in_file = rows

    def test_first_upload_no_ghosts_and_no_total_row(self, headers):
        res = _upload_chunked(self.path, "sale", headers)
        assert res.get("status") == "complete", res
        b = res.get("batch_id")
        assert b, res
        BATCHES_TO_CLEAN.append(b)
        TestIdempotencySaleUpload.b1 = b
        n_all, n_nodate, net1 = _sale_stats(b)
        TestIdempotencySaleUpload.n1 = n_all
        TestIdempotencySaleUpload.net1 = net1
        assert n_all > 0, "upload produced no records"
        assert n_nodate == 0, f"{n_nodate} ghost (date='') sale rows after upload 1"
        # numeric grand-total row must NOT be inserted -> record count should
        # not exceed rows_in_file (rows_in_file excludes the total row)
        assert n_all <= self.rows_in_file, (
            f"Inserted {n_all} > rows_in_file {self.rows_in_file} - grand-total row leaked in?"
        )

    def test_second_upload_is_idempotent(self, headers):
        res = _upload_chunked(self.path, "sale", headers)
        assert res.get("status") == "complete", res
        b2 = res.get("batch_id")
        assert b2, res
        BATCHES_TO_CLEAN.append(b2)
        n_all2, n_nodate2, net2 = _sale_stats(b2)
        assert n_nodate2 == 0, f"{n_nodate2} ghost rows after upload 2"
        assert n_all2 == self.n1, f"record count changed: {self.n1} -> {n_all2}"
        assert abs(net2 - self.net1) < 0.01, f"net_wt drift: {self.net1} -> {net2}"
        msg = (res.get("message") or "").lower()
        assert "replaced" in msg and str(self.n1) in msg, (
            f"expected 'replaced {self.n1} old records' in message, got: {res.get('message')}"
        )
        # first batch should have been fully replaced
        remaining_b1 = db.transactions.count_documents({"batch_id": self.b1})
        assert remaining_b1 == 0, f"{remaining_b1} rows from batch 1 still present"

    def test_continuation_lines_inherited_voucher_metadata(self):
        # pick any voucher with multiple lines and check they share date+refno+party
        # (use the most recent batch)
        latest_b = BATCHES_TO_CLEAN[-1]
        # find a refno that has >=2 rows
        pipe = [
            {"$match": {"batch_id": latest_b, "type": "sale"}},
            {"$group": {"_id": "$refno", "cnt": {"$sum": 1},
                        "dates": {"$addToSet": "$date"},
                        "parties": {"$addToSet": "$party_name"}}},
            {"$match": {"cnt": {"$gte": 2}}},
            {"$limit": 5},
        ]
        groups = list(db.transactions.aggregate(pipe))
        assert groups, "no multi-line vouchers found - continuation lines missing"
        for g in groups:
            assert len(g["dates"]) == 1, (
                f"refno {g['_id']} has multiple dates {g['dates']} - forward-fill failed"
            )
            assert "" not in g["dates"], f"refno {g['_id']} has empty date"
            assert len(g["parties"]) == 1, (
                f"refno {g['_id']} has multiple parties {g['parties']}"
            )

    def test_sale_return_type_inherited(self):
        latest_b = BATCHES_TO_CLEAN[-1]
        rows = list(db.transactions.find(
            {"batch_id": latest_b, "refno": self.sr_refno},
            {"_id": 0, "type": 1, "date": 1, "refno": 1},
        ))
        assert len(rows) >= 2, f"sale_return voucher {self.sr_refno} not multi-line: {rows}"
        types = {r["type"] for r in rows}
        assert types == {"sale_return"}, (
            f"continuation lines of R voucher not inherited as sale_return: {types}"
        )


class TestLegacyGhostRepair:
    """Legacy no-date sale docs must be purged by the next sale upload."""

    path = "/tmp/idempot_sale_small.xlsx"

    @classmethod
    def setup_class(cls):
        # small file
        _make_tally_sale(cls.path, n_vouchers=200, seed=99)

    def test_ghosts_are_replaced_on_next_upload(self, headers):
        ghost_batch = "legacy-ghosts-" + TAG
        # seed 150 fake legacy ghost sales
        now = datetime.datetime.utcnow().isoformat()
        docs = []
        for i in range(150):
            docs.append({
                "id": f"ghost-{TAG}-{i}",
                "type": "sale",
                "date": "",
                "refno": f"GH{i}",
                "party_name": "LEGACY GHOST",
                "item_name": "GHOST ITEM",
                "gross_wt": 1.0, "net_wt": 1.0, "fine": 0.5, "tunch": 90.0,
                "batch_id": ghost_batch,
                "upload_date": now,
            })
        db.transactions.insert_many(docs)
        assert db.transactions.count_documents({"batch_id": ghost_batch}) == 150

        res = _upload_chunked(self.path, "sale", headers)
        assert res.get("status") == "complete", res
        b = res.get("batch_id")
        BATCHES_TO_CLEAN.append(b)

        remaining_ghosts = db.transactions.count_documents({"batch_id": ghost_batch})
        assert remaining_ghosts == 0, (
            f"{remaining_ghosts}/150 legacy ghost sales still present after upload"
        )
        # no new ghosts
        new_ghosts = db.transactions.count_documents({"batch_id": b, "date": ""})
        assert new_ghosts == 0


class TestPurchaseIdempotency:
    path = "/tmp/idempot_purchase.xlsx"

    @classmethod
    def setup_class(cls):
        cls.rows = _make_small_purchase(cls.path)

    def test_purchase_uploads_are_idempotent(self, headers):
        r1 = _upload_chunked(self.path, "purchase", headers)
        assert r1.get("status") == "complete", r1
        b1 = r1.get("batch_id"); assert b1
        BATCHES_TO_CLEAN.append(b1)
        n1 = db.transactions.count_documents({"batch_id": b1})
        ghosts1 = db.transactions.count_documents({"batch_id": b1, "date": ""})
        assert n1 > 0
        assert ghosts1 == 0, f"{ghosts1} ghost purchase rows after upload 1"

        r2 = _upload_chunked(self.path, "purchase", headers)
        assert r2.get("status") == "complete", r2
        b2 = r2.get("batch_id"); assert b2
        BATCHES_TO_CLEAN.append(b2)
        n2 = db.transactions.count_documents({"batch_id": b2})
        ghosts2 = db.transactions.count_documents({"batch_id": b2, "date": ""})
        assert n2 == n1, f"purchase record count changed: {n1} -> {n2}"
        assert ghosts2 == 0
        # batch 1 fully replaced
        assert db.transactions.count_documents({"batch_id": b1}) == 0


class TestRegression:
    def test_baseline_after_cleanup(self, headers):
        # this runs before module teardown but our batches remain in DB;
        # instead validate stats endpoint & inventory endpoint still respond.
        r = requests.get(f"{BASE_URL}/api/stats", headers=headers, timeout=30)
        assert r.status_code == 200, r.text
        assert "total_transactions" in r.json()
        r2 = requests.get(f"{BASE_URL}/api/inventory/current", headers=headers, timeout=30)
        assert r2.status_code == 200

    def test_undo_upload_restores_replaced_records(self, headers):
        # Do a fresh small upload, then a second upload to trigger a replace,
        # then undo the second upload and confirm the first batch's rows are back.
        path = "/tmp/idempot_undo.xlsx"
        _make_tally_sale(path, n_vouchers=100, seed=42)
        r1 = _upload_chunked(path, "sale", headers)
        assert r1.get("status") == "complete", r1
        b1 = r1["batch_id"]; BATCHES_TO_CLEAN.append(b1)
        n_b1 = db.transactions.count_documents({"batch_id": b1})
        assert n_b1 > 0

        r2 = _upload_chunked(path, "sale", headers)
        assert r2.get("status") == "complete", r2
        b2 = r2["batch_id"]; BATCHES_TO_CLEAN.append(b2)
        # b1 should be gone
        assert db.transactions.count_documents({"batch_id": b1}) == 0

        undo = requests.post(
            f"{BASE_URL}/api/history/undo-upload",
            params={"batch_id": b2},
            headers=headers,
            timeout=60,
        )
        assert undo.status_code == 200, undo.text
        # after undo: b2 gone, b1 restored
        assert db.transactions.count_documents({"batch_id": b2}) == 0, "undo did not remove b2"
        restored = db.transactions.count_documents({"batch_id": b1})
        assert restored == n_b1, (
            f"undo did not restore replaced records: expected {n_b1}, got {restored}"
        )
