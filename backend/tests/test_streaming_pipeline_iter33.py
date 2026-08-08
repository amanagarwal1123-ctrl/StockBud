"""
Iteration 33 tests - covers:
 - Sale chunked happy path (streaming pipeline)
 - Replace + undo semantics
 - Branch-transfer chunked (>200KB) with header + OPENING BALANCE / Total skip
 - Attempts exhaustion (attempts=3, no chunks) via GET /api/upload/status
 - Unknown upload_id -> 404
 - /api/stats & /api/inventory/current regression

Auto-resume (SIGKILL) is exercised by a separate script /tmp/test_iter33_resume.py
(kept out of pytest so a hard-kill doesn't disturb the suite).
"""
import os, io, math, time, random, datetime, pytest, requests, uuid
from pymongo import MongoClient
from openpyxl import Workbook

BASE = os.environ.get("REACT_APP_BACKEND_URL", "https://sales-manager-role.preview.emergentagent.com").rstrip("/")
CHUNK = 200 * 1024
MONGO_URL = os.environ.get("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.environ.get("DB_NAME", "test_database")


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE}/api/auth/login", json={"username": "admin", "password": "admin123"}, timeout=30)
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def H(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def db():
    return MongoClient(MONGO_URL)[DB_NAME]


def _make_sale(path, n, dstart=(2025, 1, 1)):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["Some Title Row"])
    ws.append(["Date", "Type", "Refno", "Party Name", "Item Name", "Stamp", "Lbr. On Tag.No.",
               "On", "Gr.Wt.", "Gold Std.", "Fine", "Tunch", "Total", "Taxable Val.", "Pc"])
    random.seed(7)
    base = datetime.date(*dstart)
    for i in range(n):
        d = base + datetime.timedelta(days=random.randint(0, 300))
        ws.append([d, "S", f"R{i}", f"PARTY {i%50}", "CHAIN X", "S 70", "T-1",
                   round(random.uniform(5, 30), 2), round(random.uniform(0.1, 5), 3),
                   round(random.uniform(0.1, 5), 3), round(random.uniform(0.05, 4), 3),
                   round(random.uniform(60, 90), 2), round(random.uniform(1000, 9000), 2),
                   round(random.uniform(900, 8000), 2), random.randint(1, 20)])
    wb.save(path)


def _make_branch_transfer(path, n=8000):
    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.append(["Branch Transfer Book"])
    ws.append(["Date", "Type", "Refno", "Lnarr", "Gr.Wt.", "Net.Wt."])
    # First row: OPENING BALANCE (must be skipped)
    ws.append([datetime.date(2025, 1, 1), "", "OB", "OPENING BALANCE", 100.0, 100.0])
    random.seed(11)
    base = datetime.date(2025, 1, 1)
    for i in range(n):
        d = base + datetime.timedelta(days=random.randint(0, 300))
        typ = "I" if i % 2 == 0 else "R"
        ws.append([d, typ, f"BT{i}", f"CHAIN {i%20}",
                   round(random.uniform(0.5, 5), 3), round(random.uniform(0.5, 5), 3)])
    # Total row (must be skipped)
    ws.append(["", "", "", "Total", 12345.678, 12345.678])
    wb.save(path)


def _upload_file(path, H, start="2025-01-01", end="2025-12-31", file_type="sale"):
    size = os.path.getsize(path)
    total_chunks = math.ceil(size / CHUNK)
    r = requests.post(f"{BASE}/api/upload/init",
                      json={"file_type": file_type, "start_date": start, "end_date": end,
                            "total_chunks": total_chunks}, headers=H, timeout=30)
    assert r.status_code == 200, r.text
    uid = r.json()["upload_id"]
    with open(path, "rb") as f:
        for i in range(total_chunks):
            data = f.read(CHUNK)
            rr = requests.post(f"{BASE}/api/upload/chunk/{uid}?chunk_index={i}",
                               files={"file": (f"c{i}", data)}, headers=H, timeout=60)
            assert rr.status_code == 200, rr.text
    r = requests.post(f"{BASE}/api/upload/finalize/{uid}", headers=H, timeout=30)
    assert r.status_code == 200, r.text
    return uid, size, total_chunks


def _poll(uid, H, timeout_s=600):
    t0 = time.time()
    last = None
    while time.time() - t0 < timeout_s:
        r = requests.get(f"{BASE}/api/upload/status/{uid}", headers=H, timeout=30)
        if r.status_code != 200:
            return {"status": "http_error", "code": r.status_code, "body": r.text, "elapsed": time.time() - t0}
        s = r.json()
        last = s
        st = s.get("status")
        if st in ("complete", "error", "failed"):
            s["elapsed"] = time.time() - t0
            return s
        time.sleep(2)
    last["elapsed"] = time.time() - t0
    last["status"] = "timeout"
    return last


# ---------- 1. Sale chunked happy path (30K rows, 2025 dates) ----------
def test_sale_chunked_happy_path(H, db):
    # Clean 2025 data first
    db.transactions.delete_many({"date": {"$gte": "2025-01-01", "$lte": "2025-12-31"}})
    path = "/tmp/iter33_sale_30k.xlsx"
    _make_sale(path, 30000)
    size = os.path.getsize(path)
    assert size > 200 * 1024, f"file too small: {size}"
    uid, _, chunks = _upload_file(path, H)
    print(f"[sale] size={size/1024/1024:.2f}MB chunks={chunks}")
    s = _poll(uid, H, timeout_s=300)
    print(f"[sale] status={s.get('status')} elapsed={s.get('elapsed'):.1f}s msg={s.get('message')}")
    assert s["status"] == "complete", s
    batch_id = s["batch_id"]
    cnt = db.transactions.count_documents({"batch_id": batch_id})
    assert cnt == 30000, cnt
    # Chunks cleaned after complete
    ch = db.upload_chunks.count_documents({"upload_id": uid})
    assert ch == 0, f"chunks not cleaned: {ch}"
    # Cleanup
    db.transactions.delete_many({"batch_id": batch_id})
    db.replaced_records.delete_many({"batch_id": batch_id})
    db.action_history.delete_many({"batch_id": batch_id})


# ---------- 2. Replace semantics + undo ----------
def test_replace_and_undo(H, db):
    db.transactions.delete_many({"date": {"$gte": "2025-01-01", "$lte": "2025-12-31"}})
    path = "/tmp/iter33_sale_5k.xlsx"
    _make_sale(path, 5000)
    # Determine date range in file
    dates = set()
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 2:
            continue
        dt = row[0]
        if hasattr(dt, 'strftime'):
            dates.add(dt.strftime('%Y-%m-%d'))
    wb.close()
    dmin, dmax = min(dates), max(dates)
    print(f"[replace] file date range {dmin}..{dmax}")

    uid1, _, _ = _upload_file(path, H, start=dmin, end=dmax)
    s1 = _poll(uid1, H, 300); assert s1["status"] == "complete", s1
    b1 = s1["batch_id"]
    c1 = db.transactions.count_documents({"batch_id": b1})
    assert c1 == 5000, c1

    uid2, _, _ = _upload_file(path, H, start=dmin, end=dmax)
    s2 = _poll(uid2, H, 300); assert s2["status"] == "complete", s2
    b2 = s2["batch_id"]
    print(f"[replace] s2 msg={s2.get('message')} replaced={s2.get('replaced')}")

    parts = list(db.replaced_records.find({"batch_id": b2}))
    assert len(parts) >= 1, "no replaced_records parts"
    for p in parts:
        assert "part" in p

    # Undo second batch -> restore first
    r = requests.post(f"{BASE}/api/history/undo-upload?batch_id={b2}", headers=H, timeout=120)
    assert r.status_code == 200, r.text
    restored = db.transactions.count_documents({"type": "sale",
                                                 "date": {"$gte": dmin, "$lte": dmax}})
    assert restored == 5000, restored

    # Cleanup all 2025
    db.transactions.delete_many({"date": {"$gte": "2025-01-01", "$lte": "2025-12-31"}})
    db.replaced_records.delete_many({"batch_id": {"$in": [b1, b2]}})
    db.action_history.delete_many({"batch_id": {"$in": [b1, b2]}})


# ---------- 3. Branch transfer chunked (NEW - was broken before) ----------
def test_branch_transfer_chunked(H, db):
    db.transactions.delete_many({"type": {"$in": ["issue", "receive"]},
                                 "date": {"$gte": "2025-01-01", "$lte": "2025-12-31"}})
    path = "/tmp/iter33_bt.xlsx"
    _make_branch_transfer(path, n=8000)
    size = os.path.getsize(path)
    print(f"[bt] size={size/1024:.1f}KB")
    assert size > 200 * 1024, f"file too small ({size})"
    uid, _, chunks = _upload_file(path, H, file_type="branch_transfer")
    s = _poll(uid, H, 300)
    print(f"[bt] status={s.get('status')} elapsed={s.get('elapsed'):.1f}s msg={s.get('message')}")
    assert s["status"] == "complete", s
    batch_id = s["batch_id"]
    cnt = db.transactions.count_documents({"batch_id": batch_id})
    # OPENING BALANCE + Total both skipped, 8000 data rows survive
    assert cnt == 8000, cnt
    # verify types
    n_issue = db.transactions.count_documents({"batch_id": batch_id, "type": "issue"})
    n_recv = db.transactions.count_documents({"batch_id": batch_id, "type": "receive"})
    print(f"[bt] issue={n_issue} receive={n_recv}")
    assert n_issue > 0 and n_recv > 0
    assert n_issue + n_recv == 8000
    # party_name should be MMI Jewelly Branch
    sample = db.transactions.find_one({"batch_id": batch_id})
    print(f"[bt] sample party={sample.get('party_name')} item={sample.get('item_name')}")
    assert sample.get("party_name") == "MMI Jewelly Branch"
    # No OPENING BALANCE / Total leaked
    assert db.transactions.count_documents({"batch_id": batch_id,
                                            "item_name": {"$regex": "^(OPENING BALANCE|Total)$", "$options": "i"}}) == 0
    # Cleanup
    db.transactions.delete_many({"batch_id": batch_id})
    db.replaced_records.delete_many({"batch_id": batch_id})
    db.action_history.delete_many({"batch_id": batch_id})


# ---------- 4. Attempts exhaustion ----------
def test_attempts_exhausted(H, db):
    fake_uid = f"iter33_exhaust_{uuid.uuid4().hex[:8]}"
    stale = (datetime.datetime.now(datetime.timezone.utc) - datetime.timedelta(minutes=10)).isoformat()
    db.upload_sessions.insert_one({
        "upload_id": fake_uid,
        "status": "processing",
        "owner_username": "admin",
        "attempts": 3,
        "heartbeat": stale,
        "file_type": "sale",
        "start_date": "2025-01-01",
        "end_date": "2025-12-31",
        "created_at": stale,
    })
    try:
        r = requests.get(f"{BASE}/api/upload/status/{fake_uid}", headers=H, timeout=30)
        assert r.status_code == 200, r.text
        s = r.json()
        print(f"[exhaust] status={s.get('status')} detail={s.get('detail')}")
        assert s["status"] == "error", s
        blob = ((s.get("detail") or "") + " " + (s.get("message") or "")).lower()
        assert "retry" in blob or "exhaust" in blob or "attempt" in blob or "re-upload" in blob, s
    finally:
        db.upload_sessions.delete_many({"upload_id": fake_uid})


# ---------- 5. Unknown upload_id 404 ----------
def test_unknown_upload_id_404(H):
    r = requests.get(f"{BASE}/api/upload/status/does_not_exist_zzz", headers=H, timeout=30)
    assert r.status_code == 404, r.text


# ---------- 6. Regression: /api/stats + /api/inventory/current ----------
def test_regression_stats_and_inventory(H, db):
    # Ensure no leftover 2025 rows (defensive)
    db.transactions.delete_many({"date": {"$gte": "2025-01-01", "$lte": "2025-12-31"}})
    r = requests.get(f"{BASE}/api/stats", headers=H, timeout=30)
    assert r.status_code == 200, r.text
    s = r.json()
    print(f"[reg] total_transactions={s.get('total_transactions')}")
    assert s.get("total_transactions") == 11767, s

    r = requests.get(f"{BASE}/api/inventory/current", headers=H, timeout=60)
    assert r.status_code == 200, r.text
